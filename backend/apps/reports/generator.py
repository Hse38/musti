from io import BytesIO

import openpyxl
from openpyxl.styles import Font, PatternFill

from apps.competitions.models import Competition, Participant
from apps.invoices.models import Invoice
from apps.transport.models import TransportRequest


RESULT_COLUMNS = [
    {"key": "order", "label": "Sıra", "order": 1},
    {"key": "competition", "label": "Yarışma Adı", "order": 2},
    {"key": "first_name", "label": "Ad", "order": 3},
    {"key": "last_name", "label": "Soyad", "order": 4},
    {"key": "email", "label": "Email", "order": 5},
    {"key": "phone", "label": "Telefon Numarası", "order": 6},
    {"key": "tc_id", "label": "TC", "order": 7},
    {"key": "transport_type", "label": "Ulaşım tipi", "order": 8},
    {"key": "invoice_status", "label": "Fatura durumu", "order": 9},
    {"key": "errors", "label": "Tespit edilen hatalar", "order": 10},
    {"key": "decision", "label": "Nihai karar (Onay / Red)", "order": 11},
    {"key": "rejection_reason", "label": "Red nedeni (varsa)", "order": 12},
]

PAYMENT_COLUMNS = [
    {"key": "order", "label": "Sayı", "order": 1},
    {"key": "account_holder_name", "label": "Hesap Sahibinin Adı-Soyadı", "order": 2},
    {
        "key": "account_holder_tc",
        "label": "Hesap Sahibinin T.C. Kimlik Numarası",
        "order": 3,
    },
    {"key": "bank_name", "label": "Banka Adı", "order": 4},
    {"key": "iban", "label": "IBAN", "order": 5},
    {"key": "amount", "label": "Tutar", "order": 6},
    {"key": "description", "label": "Açıklama", "order": 7},
]


class ReportGenerator:
    GREEN = PatternFill("solid", fgColor="C6EFCE")
    RED = PatternFill("solid", fgColor="FFC7CE")
    BLUE = PatternFill("solid", fgColor="BDD7EE")
    YELLOW = PatternFill("solid", fgColor="FFEB9C")

    def generate_result_report(self, session, template=None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sonuç Raporu"

        columns = self._get_columns("result", template)
        ws.append([col["label"] for col in columns])
        self._style_header(ws)

        subs = list(
            session.submissions.select_related("participant__team").order_by("id")
        )
        comp_name = session.competition.name if session.competition else ""

        for idx, sub in enumerate(subs, start=1):
            row = self._build_result_row(sub, columns, idx, comp_name)
            ws.append(row)
            fill = self.GREEN if sub.status == "approved" else self.RED
            for cell in ws[ws.max_row]:
                cell.fill = fill

        return self._to_bytes(wb)

    def generate_payment_report(self, session, template=None) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ödeme Raporu"

        approved = session.submissions.filter(status="approved").select_related(
            "participant", "session__competition"
        )
        columns = self._get_columns("payment", template)
        ws.append([col["label"] for col in columns])
        self._style_header(ws)

        comp = session.competition
        comp_name = comp.name if comp else "TEKNOFEST"
        year = comp.start_date.year if comp and comp.start_date else 2025

        for idx, sub in enumerate(approved.order_by("id"), start=1):
            p = sub.participant
            snap = sub.payment_form_snapshot or {}
            row_data = {
                "order": idx,
                "account_holder_name": (
                    p.account_holder_name
                    if p
                    else snap.get("account_holder_name", "")
                ),
                "account_holder_tc": sub.account_holder_tc
                or (p.tc_id if p else snap.get("account_holder_tc", "")),
                "bank_name": (p.bank_name if p else snap.get("bank_name", "")),
                "iban": (p.iban if p else snap.get("iban", "")),
                "amount": float(sub.invoice_amount or 0),
                "description": f"TEKNOFEST {year} {comp_name} BİLET ÖDEMESİ",
            }
            ws.append([row_data.get(col["key"], "") for col in columns])

        return self._to_bytes(wb)

    def generate_flight_report(self) -> bytes:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Uçak"
        headers = [
            "Ad Soyad",
            "TC",
            "Takım ID",
            "Kalkış / menşei",
            "Varış tarihi",
            "Dönüş tarihi",
            "Uçuş notları",
            "Durum",
        ]
        ws.append(headers)
        self._style_header(ws)
        qs = TransportRequest.objects.filter(transport_type="plane").select_related(
            "participant",
            "participant__team",
        )
        for tr in qs:
            p = tr.participant
            team_id = ""
            if p.team:
                team_id = p.team.team_id or p.team.team_code or ""
            ws.append(
                [
                    p.full_name,
                    p.tc_id or "",
                    team_id,
                    tr.origin_city,
                    tr.preferred_arrival_date or "",
                    tr.preferred_departure_date or "",
                    tr.flight_notes,
                    tr.get_status_display(),
                ]
            )
        return self._to_bytes(wb)

    def generate_result_report_competition(self, competition_id: int) -> bytes:
        comp = Competition.objects.get(pk=competition_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Sonuç"
        headers = [
            "Sıra",
            "Yarışma Adı",
            "Ad",
            "Soyad",
            "Email",
            "Telefon",
            "TC",
            "Ulaşım Tipi",
            "Fatura Durumu",
            "Tespit Edilen Hatalar",
            "Nihai Karar",
            "Red Nedeni",
        ]
        ws.append(headers)
        self._style_header(ws)
        qs = Participant.objects.filter(team__competition=comp).select_related("team").order_by(
            "id"
        )
        for idx, p in enumerate(qs, start=1):
            parts = (p.full_name or "").split()
            first = parts[0] if parts else ""
            last = " ".join(parts[1:]) if len(parts) > 1 else ""
            tr = TransportRequest.objects.filter(participant=p).first()
            inv = None
            if tr:
                inv = tr.invoices.order_by("-id").first()
            inv_status = inv.get_status_display() if inv else "—"
            errs = ""
            decision = "Bekliyor"
            reason = ""
            fill = self.YELLOW
            if inv:
                errs = "\n".join(inv.rejection_reasons or []) if inv.rejection_reasons else ""
                if inv.status == "approved":
                    decision = "Onay"
                    fill = self.GREEN
                elif inv.status == "rejected":
                    decision = "Red"
                    reason = errs
                    fill = self.RED
                elif inv.status == "manual_review":
                    decision = "İncelemede"
                    fill = self.YELLOW
            if p.transport_type == "plane":
                fill = self.BLUE
            row = [
                idx,
                comp.name,
                first,
                last,
                p.email or "",
                p.phone or "",
                p.tc_id,
                p.get_transport_type_display(),
                inv_status,
                errs,
                decision,
                reason,
            ]
            ws.append(row)
            for cell in ws[ws.max_row]:
                cell.fill = fill
        return self._to_bytes(wb)

    def generate_payment_report_competition(self, competition_id: int) -> bytes:
        comp = Competition.objects.get(pk=competition_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Ödeme"
        headers = [
            "Sayı",
            "Hesap Sahibinin Adı-Soyadı",
            "Hesap Sahibinin T.C. Kimlik Numarası",
            "Banka Adı",
            "IBAN",
            "Tutar",
            "Açıklama",
        ]
        ws.append(headers)
        self._style_header(ws)
        year = comp.year or (comp.start_date.year if comp.start_date else 2025)
        desc = f"TEKNOFEST {year} {comp.name} BİLET ÖDEMESİ"
        invs = (
            Invoice.objects.filter(
                status="approved",
                transport_request__participant__team__competition=comp,
            )
            .select_related("transport_request__participant")
            .order_by("id")
        )
        for idx, inv in enumerate(invs, start=1):
            p = inv.transport_request.participant
            ws.append(
                [
                    idx,
                    p.account_holder_name or p.full_name,
                    p.tc_id,
                    p.bank_name,
                    p.iban,
                    float(inv.ai_extracted_amount or 0),
                    desc,
                ]
            )
        return self._to_bytes(wb)

    def generate_flights_report_competition(self, competition_id: int) -> bytes:
        comp = Competition.objects.get(pk=competition_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Uçak"
        headers = [
            "Sıra",
            "Ad Soyad",
            "TC",
            "Email",
            "Telefon",
            "Takım",
            "Tercih Geliş Tarihi",
            "Tercih Dönüş Tarihi",
            "Notlar",
        ]
        ws.append(headers)
        self._style_header(ws)
        qs = (
            TransportRequest.objects.filter(
                transport_type="plane",
                participant__team__competition=comp,
            )
            .select_related("participant", "participant__team")
            .order_by("id")
        )
        for idx, tr in enumerate(qs, start=1):
            p = tr.participant
            ws.append(
                [
                    idx,
                    p.full_name,
                    p.tc_id,
                    p.email or "",
                    p.phone or "",
                    p.team.name if p.team else "",
                    str(tr.preferred_arrival_date or ""),
                    str(tr.preferred_departure_date or ""),
                    tr.flight_notes or "",
                ]
            )
        return self._to_bytes(wb)

    def generate_tracking_report_competition(self, competition_id: int) -> bytes:
        comp = Competition.objects.get(pk=competition_id)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Takip"
        headers = [
            "Ad Soyad",
            "TC",
            "Takım",
            "Yarışma",
            "Magic Link Gönderildi mi",
            "Gönderilme Tarihi",
            "İlk Giriş Tarihi",
            "Ulaşım Tercihi",
            "Tercih Tarihi",
            "Fatura Yükleme Tarihi",
            "Son Durum",
            "Son Güncelleme",
        ]
        ws.append(headers)
        self._style_header(ws)
        qs = Participant.objects.filter(team__competition=comp).select_related("team")
        for p in qs.order_by("id"):
            tr = TransportRequest.objects.filter(participant=p).first()
            last_inv = None
            if tr:
                last_inv = tr.invoices.order_by("-id").first()
            status_txt = "—"
            if last_inv:
                status_txt = last_inv.get_status_display()
            elif tr and tr.invoices.exists():
                status_txt = "Fatura var"
            elif tr and tr.status != "pending":
                status_txt = tr.get_status_display()
            ws.append(
                [
                    p.full_name,
                    p.tc_id,
                    p.team.name if p.team else "",
                    comp.name,
                    "Evet" if p.magic_link_sent_at else "Hayır",
                    p.magic_link_sent_at.isoformat() if p.magic_link_sent_at else "",
                    p.first_login_at.isoformat() if p.first_login_at else "",
                    p.get_transport_type_display(),
                    p.transport_selected_at.isoformat() if p.transport_selected_at else "",
                    p.invoice_uploaded_at.isoformat() if p.invoice_uploaded_at else "",
                    status_txt,
                    p.last_activity_at.isoformat() if p.last_activity_at else "",
                ]
            )
        return self._to_bytes(wb)

    def _build_result_row(self, sub, columns, order_idx: int, competition_name: str):
        p = sub.participant
        snap = sub.payment_form_snapshot or {}
        full = (
            p.full_name
            if p
            else (snap.get("full_name") or "")
        )
        parts = full.split()
        first_name = parts[0] if parts else ""
        last_name = " ".join(parts[1:]) if len(parts) > 1 else ""

        if p:
            transport = p.get_transport_type_display()
            email = p.email or ""
            phone = p.phone or ""
            tc = p.tc_id
        else:
            transport = snap.get("transport_type_declared", "") or ""
            email = snap.get("email", "")
            phone = snap.get("phone", "")
            tc = snap.get("tc_id", "")

        has_inv = bool(sub.invoice_file or sub.invoice_drive_link)
        invoice_status = "Var" if has_inv else "Yok"
        errs = "\n".join(sub.rejection_reasons or [])
        decision = "Onay" if sub.status == "approved" else "Red"
        rejection_reason = errs if sub.status == "rejected" else ""

        base = {
            "order": order_idx,
            "competition": competition_name,
            "first_name": first_name,
            "last_name": last_name,
            "email": email,
            "phone": phone,
            "tc_id": tc,
            "transport_type": transport,
            "invoice_status": invoice_status,
            "errors": errs,
            "decision": decision,
            "rejection_reason": rejection_reason,
        }
        return [base.get(col["key"], "") for col in columns]

    def _get_columns(self, report_type, template):
        if template and template.columns:
            cols = template.columns
            if isinstance(cols, list) and cols and isinstance(cols[0], dict):
                if "order" in cols[0]:
                    return sorted(cols, key=lambda x: x.get("order", 0))
            return cols
        defaults = {
            "result": RESULT_COLUMNS,
            "payment": PAYMENT_COLUMNS,
        }
        return sorted(defaults[report_type], key=lambda x: x["order"])

    def _style_header(self, ws):
        bold = Font(bold=True)
        gray = PatternFill("solid", fgColor="D9D9D9")
        for cell in ws[1]:
            cell.font = bold
            cell.fill = gray

    def _to_bytes(self, wb) -> bytes:
        buf = BytesIO()
        wb.save(buf)
        return buf.getvalue()
